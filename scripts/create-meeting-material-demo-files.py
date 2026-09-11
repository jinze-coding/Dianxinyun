#!/usr/bin/env python3
"""Generate synthetic meeting material fixtures; never reads real project documents."""
from pathlib import Path
import subprocess
import zipfile
from docx import Document
from pptx import Presentation
from openpyxl import Workbook
from PIL import Image, ImageDraw

root = Path(__file__).resolve().parent.parent / "logs/meeting-material-demo-files"
root.mkdir(parents=True, exist_ok=True)
document = Document()
document.add_heading("会议通知 · 开发演示", 0)
document.add_paragraph("本资料为合成测试数据，不代表真实会议安排。")
document.add_heading("参会说明", 1)
document.add_paragraph("请提前填写预约，到场扫描会场签到码逐人签到。")
document.add_paragraph("议程：质量检查、安全交底、问题讨论和任务确认。")
document.save(root / "会议通知.docx")
workbook = Workbook()
sheet = workbook.active
sheet.title = "演示议程"
for row in [["时间", "议题", "备注"], ["09:00", "质量检查", "演示数据"], ["09:30", "安全交底", "演示数据"]]:
    sheet.append(row)
sheet.column_dimensions["B"].width = 26
workbook.save(root / "会议议程.xlsx")
slides = Presentation()
slide = slides.slides.add_slide(slides.slide_layouts[0])
slide.shapes.title.text = "质量安全周例会"
slide.placeholders[1].text = "开发演示课件\n预约、会场签到与会议资料留档"
slides.save(root / "会议课件.pptx")
image = Image.new("RGB", (1200, 700), "#edf4ff")
draw = ImageDraw.Draw(image)
draw.rounded_rectangle((60, 60, 1140, 640), radius=28, fill="white", outline="#1677ff", width=3)
draw.text((110, 130), "MEETING MATERIALS / LOCAL DEMO", fill="#1677ff", font_size=42)
draw.text((110, 240), "Registration  >  Check-in  >  Archive", fill="#24415b", font_size=32)
draw.text((110, 430), "Synthetic sample. No personal information.", fill="#6a7f95", font_size=24)
image.save(root / "现场示意.png")
(root / "内部会议纪要.txt").write_text("内部资料 · 合成测试\n请确认责任分工与整改时限。\n仅用于开发验证，无真实人员信息。", encoding="utf-8")
(root / "公开参会说明.txt").write_text("公开资料 · 合成测试\n会后可通过原邀请码继续查看会议公开资料。", encoding="utf-8")
(root / "会议纪要修订版.txt").write_text("内部修订草稿 V2\n上传新版本不会自动公开。需要管理人员明确发布。", encoding="utf-8")
(root / "演示图纸.dxf").write_text("0\nSECTION\n2\nHEADER\n0\nENDSEC\n0\nEOF\n")
with zipfile.ZipFile(root / "会议归档包.zip", "w") as archive:
    archive.writestr("README.txt", "Synthetic meeting archive for local development.")
subprocess.run(["ffmpeg", "-nostdin", "-v", "error", "-f", "lavfi", "-i", "color=c=0x1677ff:s=640x360:d=3", "-c:v", "libx264", "-y", str(root / "现场视频.mp4")], check=True)
subprocess.run(["ffmpeg", "-nostdin", "-v", "error", "-f", "lavfi", "-i", "sine=frequency=440:duration=2", "-c:a", "libmp3lame", "-y", str(root / "提示音.mp3")], check=True)
print(root)
